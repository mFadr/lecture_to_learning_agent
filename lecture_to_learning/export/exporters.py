"""Export výstupů: Markdown, HTML, DOCX, CSV/TSV (Anki), JSON, PNG/SVG."""

from __future__ import annotations

import base64
import csv
import io
import json
import logging
import shutil
from html import escape
from pathlib import Path

from ..models import Origin, Project, StudyResults, SummaryPoint, Visualization, format_pages
from ..security import ensure_within, safe_filename, unique_path

log = logging.getLogger(__name__)

ORIGIN_MARK = {Origin.SOURCE: "", Origin.DERIVED: "", Origin.MODEL: " *(externí vysvětlení – není v PDF)*"}
ORIGIN_HTML = {Origin.SOURCE: "", Origin.DERIVED: "",
               Origin.MODEL: ' <span class="ext">externí vysvětlení – není v PDF</span>'}


def _src(project: Project, doc_id: str, pages: list[int]) -> str:
    name = project.doc_names().get(doc_id, "?")
    return f"{name}, s. {format_pages(pages)}" if pages else f"{name}, stránka neurčena"


# ---------------------------------------------------------------- Markdown


def summary_markdown(project: Project, r: StudyResults) -> str:
    out = [f"# {project.name} – studijní shrnutí", ""]
    if r.summary:
        out.append(f"*Úroveň detailu: {r.summary.detail.label}. Body označené „externí vysvětlení“ nejsou obsahem "
                   "přednášky, ale doplňkem jazykového modelu.*\n")
    if not r.summary or not r.summary.sections:
        out.append("_Shrnutí zatím nebylo vytvořeno._")
    else:
        for s in r.summary.sections:
            out += [f"## {s.title}", f"*Zdroj: {_src(project, s.source.doc_id, s.source.pages)}*", ""]
            if s.overview:
                out += [s.overview, ""]
            for title, pts in (("Vysvětlení", s.explanation), ("Klíčové definice", s.key_definitions),
                               ("Vzorce", s.formulas), ("Příklady z přednášky", s.examples),
                               ("Souvislosti", s.connections), ("Časté chyby a nejasnosti", s.pitfalls)):
                if pts:
                    out.append(f"### {title}")
                    out += [_md_point(p) for p in pts]
                    out.append("")
            if s.missing_info:
                out.append("> **Ve zdroji chybí / nejasné:** " + "; ".join(s.missing_info))
                out.append("")
    if r.terms:
        out += ["---", "## Klíčové pojmy a definice", ""]
        out += terms_markdown_lines(project, r)
    return "\n".join(out)


def _md_point(p: SummaryPoint) -> str:
    pages = f" *(s. {format_pages(p.pages)})*" if p.pages else ""
    return f"- {p.text}{pages}{ORIGIN_MARK[p.origin]}"


def terms_markdown_lines(project: Project, r: StudyResults) -> list[str]:
    out = []
    for t in r.terms:
        out.append(f"**{t.name}** — {t.definition}  ")
        if t.source.quote_verified and t.source.quote:
            out.append(f"> „{t.source.quote}“ (doslovná citace)  ")
        if t.simple_explanation:
            out.append(f"*Jednoduše:* {t.simple_explanation}  ")
        if t.example:
            out.append(f"*Příklad:* {t.example}  ")
        if t.related:
            out.append(f"*Související:* {', '.join(t.related)}  ")
        out.append(f"*Zdroj: {_src(project, t.source.doc_id, t.source.pages)}*\n")
    return out


def quiz_markdown(project: Project, r: StudyResults, with_answers: bool = True) -> str:
    out = [f"# {project.name} – procvičovací test", ""]
    for i, q in enumerate(r.quiz, 1):
        out.append(f"**{i}. ({q.kind.label})** {q.question}")
        for j, o in enumerate(q.options):
            out.append(f"   {'ABCDEF'[j]}) {o}")
        out.append("")
    if with_answers:
        out += ["---", "## Řešení", ""]
        for i, q in enumerate(r.quiz, 1):
            out.append(f"**{i}.** {q.correct_answer}  ")
            if q.explanation:
                out.append(f"{q.explanation}  ")
            out.append(f"*Zdroj: {_src(project, q.source.doc_id, q.source.pages)}*\n")
    return "\n".join(out)


# -------------------------------------------------------------------- HTML

_CSS = """body{font-family:'Segoe UI',Arial,sans-serif;max-width:920px;margin:2em auto;padding:0 1em;color:#1f2933;
line-height:1.55}h1{color:#1e3a8a}h2{border-bottom:2px solid #dbeafe;padding-bottom:.2em;margin-top:2em}
.src{color:#52606d;font-size:.9em}.ext{background:#fef3c7;color:#92400e;font-size:.8em;padding:0 .4em;border-radius:3px}
.missing{background:#fef2f2;border-left:4px solid #b91c1c;padding:.5em 1em}.term{margin:.8em 0}
blockquote{border-left:3px solid #93c5fd;margin:.3em 0;padding-left:.8em;color:#334e68}
figure{margin:1.5em 0}figure img{max-width:100%;border:1px solid #e4e7eb}figcaption{font-size:.92em}
.q{margin:1em 0}.opt{margin-left:1.5em}details{margin:.3em 0 1em 1.5em;color:#334e68}"""


def _h(s: str) -> str:
    return escape(s, quote=True)


def _html_point(p: SummaryPoint) -> str:
    pages = f' <span class="src">(s. {_h(format_pages(p.pages))})</span>' if p.pages else ""
    return f"<li>{_h(p.text)}{pages}{ORIGIN_HTML[p.origin]}</li>"


def _img_data_uri(path: Path) -> str | None:
    if not path.exists():
        return None
    mime = "image/svg+xml" if path.suffix.lower() == ".svg" else "image/png"
    return f"data:{mime};base64," + base64.b64encode(path.read_bytes()).decode("ascii")


def study_html(project: Project, r: StudyResults, project_dir: Path, include_quiz: bool = True) -> str:
    """Samostatný HTML dokument. Veškerý text je escapován; obrázky jsou vloženy jako <img> (bez skriptů)."""
    parts = [f"<!DOCTYPE html><html lang='cs'><head><meta charset='utf-8'><title>{_h(project.name)}</title>"
             f"<meta http-equiv='Content-Security-Policy' content=\"default-src 'none'; img-src data:; style-src 'unsafe-inline'\">"
             f"<style>{_CSS}</style></head><body>", f"<h1>{_h(project.name)}</h1>",
             "<p class='src'>Zdroje: " + _h(", ".join(d.filename for d in project.documents)) + "</p>"]
    if r.summary and r.summary.sections:
        parts.append(f"<p><em>{_h(r.summary.detail.label)}.</em> Body označené "
                     "<span class='ext'>externí vysvětlení</span> nejsou obsahem přednášky.</p>")
        for s in r.summary.sections:
            parts.append(f"<h2>{_h(s.title)}</h2><p class='src'>Zdroj: {_h(_src(project, s.source.doc_id, s.source.pages))}</p>")
            if s.overview:
                parts.append(f"<p>{_h(s.overview)}</p>")
            for title, pts in (("Vysvětlení", s.explanation), ("Klíčové definice", s.key_definitions),
                               ("Vzorce", s.formulas), ("Příklady z přednášky", s.examples),
                               ("Souvislosti", s.connections), ("Časté chyby a nejasnosti", s.pitfalls)):
                if pts:
                    parts.append(f"<h3>{title}</h3><ul>" + "".join(_html_point(p) for p in pts) + "</ul>")
            if s.missing_info:
                parts.append("<div class='missing'><strong>Ve zdroji chybí / nejasné:</strong> "
                             + _h("; ".join(s.missing_info)) + "</div>")
    if r.terms:
        parts.append("<h2>Klíčové pojmy a definice</h2>")
        for t in r.terms:
            q = (f"<blockquote>„{_h(t.source.quote)}“ <span class='src'>(doslovná citace)</span></blockquote>"
                 if t.source.quote_verified and t.source.quote else "")
            parts.append(f"<div class='term'><strong>{_h(t.name)}</strong> — {_h(t.definition)}{q}"
                         + (f"<br><em>Jednoduše:</em> {_h(t.simple_explanation)}" if t.simple_explanation else "")
                         + (f"<br><em>Příklad:</em> {_h(t.example)}" if t.example else "")
                         + f"<br><span class='src'>{_h(_src(project, t.source.doc_id, t.source.pages))}</span></div>")
    valid = latest_visualizations(r)
    if valid or r.topic_map_svg:
        parts.append("<h2>Vizualizace</h2>")
        if r.topic_map_svg:
            uri = _img_data_uri(project_dir / r.topic_map_svg)
            if uri:
                parts.append(f"<figure><img src='{uri}' alt='Tematická mapa'><figcaption>Tematická mapa</figcaption></figure>")
        for v in valid:
            path = project_dir / (v.svg_path or "")
            uri = _img_data_uri(path) if v.svg_path else None
            if uri:
                parts.append(f"<figure><img src='{uri}' alt='{_h(v.spec.title)}'><figcaption><strong>{_h(v.spec.title)}"
                             f"</strong> – {_h(v.spec.explanation)}<br><span class='src'>Zdroj: "
                             f"{_h(_src(project, v.source.doc_id, v.source.pages))}</span></figcaption></figure>")
    if include_quiz and r.quiz:
        parts.append("<h2>Procvičovací test</h2>")
        for i, q in enumerate(r.quiz, 1):
            opts = "".join(f"<div class='opt'>{'ABCDEF'[j]}) {_h(o)}</div>" for j, o in enumerate(q.options))
            parts.append(f"<div class='q'><strong>{i}.</strong> <em>({_h(q.kind.label)})</em> {_h(q.question)}{opts}"
                         f"<details><summary>Řešení</summary><strong>{_h(q.correct_answer)}</strong><br>"
                         f"{_h(q.explanation)}<br><span class='src'>{_h(_src(project, q.source.doc_id, q.source.pages))}"
                         "</span></details></div>")
    parts.append("</body></html>")
    return "\n".join(parts)


def latest_visualizations(r: StudyResults) -> list[Visualization]:
    """Poslední platná verze vizualizace pro každý příklad."""
    best: dict[str, Visualization] = {}
    for v in r.visualizations:
        if not v.valid:
            continue
        k = v.example_id or v.id
        if k not in best or v.version > best[k].version:
            best[k] = v
    return list(best.values())


# -------------------------------------------------------------------- DOCX


def study_docx(project: Project, r: StudyResults, project_dir: Path, out: Path) -> Path:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor

    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_heading(project.name, 0)
    doc.add_paragraph("Zdroje: " + ", ".join(d.filename for d in project.documents)).italic = True

    def point(p: SummaryPoint) -> None:
        para = doc.add_paragraph(style="List Bullet")
        para.add_run(p.text)
        if p.pages:
            run = para.add_run(f" (s. {format_pages(p.pages)})")
            run.font.color.rgb = RGBColor(0x52, 0x60, 0x6D)
        if p.origin == Origin.MODEL:
            run = para.add_run(" [externí vysvětlení – není v PDF]")
            run.font.color.rgb = RGBColor(0x92, 0x40, 0x0E)
            run.italic = True

    if r.summary:
        for s in r.summary.sections:
            doc.add_heading(s.title, 1)
            src = doc.add_paragraph(f"Zdroj: {_src(project, s.source.doc_id, s.source.pages)}")
            src.runs[0].font.color.rgb = RGBColor(0x52, 0x60, 0x6D)
            if s.overview:
                doc.add_paragraph(s.overview)
            for title, pts in (("Vysvětlení", s.explanation), ("Klíčové definice", s.key_definitions),
                               ("Vzorce", s.formulas), ("Příklady z přednášky", s.examples),
                               ("Souvislosti", s.connections), ("Časté chyby a nejasnosti", s.pitfalls)):
                if pts:
                    doc.add_heading(title, 2)
                    for p in pts:
                        point(p)
            if s.missing_info:
                doc.add_paragraph("Ve zdroji chybí / nejasné: " + "; ".join(s.missing_info)).italic = True
    if r.terms:
        doc.add_heading("Klíčové pojmy a definice", 1)
        for t in r.terms:
            para = doc.add_paragraph()
            para.add_run(t.name).bold = True
            para.add_run(f" — {t.definition}")
            if t.simple_explanation:
                para.add_run(f"\nJednoduše: {t.simple_explanation}").italic = True
            para.add_run(f"\n{_src(project, t.source.doc_id, t.source.pages)}").font.color.rgb = RGBColor(0x52, 0x60, 0x6D)
    vis = [v for v in latest_visualizations(r) if v.png_path and (project_dir / v.png_path).exists()]
    if vis:
        doc.add_heading("Vizualizace", 1)
        for v in vis:
            doc.add_heading(v.spec.title, 2)
            try:
                doc.add_picture(str(project_dir / v.png_path), width=Inches(6))
            except Exception as exc:
                log.warning("Obrázek nelze vložit do DOCX: %s", exc)
            doc.add_paragraph(v.spec.explanation)
            doc.add_paragraph(f"Zdroj: {_src(project, v.source.doc_id, v.source.pages)}").italic = True
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    return out


# ------------------------------------------------------------- Flashcards


def flashcards_csv(project: Project, r: StudyResults) -> str:
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";", quoting=csv.QUOTE_MINIMAL, lineterminator="\r\n")
    w.writerow(["otazka", "odpoved", "typ", "tema", "zdroj"])
    for c in r.flashcards:
        w.writerow([c.question, c.answer, c.kind.value, c.topic, _src(project, c.source.doc_id, c.source.pages)])
    return buf.getvalue()


def _anki_field(s: str) -> str:
    return s.replace("\t", " ").replace("\r", "").replace("\n", " ").strip()


def _anki_tag(s: str) -> str:
    return "_".join(s.split())[:60]


def flashcards_anki_tsv(project: Project, r: StudyResults) -> str:
    """Formát pro import do Anki (Soubor → Importovat): přední strana, zadní strana, štítky."""
    lines = ["#separator:tab", "#html:false", "#tags column:3"]
    deck_tag = _anki_tag(project.name)
    for c in r.flashcards:
        back = f"{c.answer} ({_src(project, c.source.doc_id, c.source.pages)})"
        tags = " ".join(t for t in (deck_tag, _anki_tag(c.topic), c.kind.value) if t)
        lines.append("\t".join([_anki_field(c.question), _anki_field(back), tags]))
    return "\n".join(lines) + "\n"


# ----------------------------------------------------------- souhrnný export


def export_all(project: Project, r: StudyResults, project_dir: Path, out_dir: Path,
               formats: set[str] | None = None) -> list[Path]:
    """Exportuje vybrané formáty do složky `out_dir` (vytvoří podsložku s názvem projektu)."""
    formats = formats or {"md", "html", "docx", "csv", "tsv", "json", "png", "svg"}
    target = unique_path(out_dir / safe_filename(f"{project.name} – export", has_ext=False))
    target.mkdir(parents=True, exist_ok=False)
    written: list[Path] = []

    def write(name: str, text: str, encoding: str = "utf-8") -> None:
        p = ensure_within(target, Path(safe_filename(name)))
        p.write_text(text, encoding=encoding, newline="")
        written.append(p)

    if "md" in formats:
        write("shrnuti.md", summary_markdown(project, r))
        if r.quiz:
            write("test.md", quiz_markdown(project, r))
    if "html" in formats:
        write("studijni_material.html", study_html(project, r, project_dir))
    if "docx" in formats:
        written.append(study_docx(project, r, project_dir, target / "studijni_material.docx"))
    if "csv" in formats and r.flashcards:
        write("flashcards.csv", flashcards_csv(project, r), encoding="utf-8-sig")  # BOM kvůli Excelu
    if "tsv" in formats and r.flashcards:
        write("flashcards_anki.txt", flashcards_anki_tsv(project, r))
    if "json" in formats:
        write("vysledky.json", json.dumps(r.model_dump(mode="json"), ensure_ascii=False, indent=1))
    if formats & {"png", "svg"}:
        vdir = target / "vizualizace"
        vdir.mkdir()
        for v in latest_visualizations(r):
            for kind in ("png", "svg"):
                rel = v.png_path if kind == "png" else v.svg_path
                if kind in formats and rel and (project_dir / rel).exists():
                    dst = unique_path(vdir / safe_filename(f"{v.spec.title}.{kind}"))
                    shutil.copy2(project_dir / rel, dst)
                    written.append(dst)
            if v.mermaid:
                dst = unique_path(vdir / safe_filename(f"{v.spec.title}.mmd"))
                dst.write_text(v.mermaid, encoding="utf-8")
                written.append(dst)
        for rel in (r.topic_map_svg, r.relation_map_svg):
            if rel and (project_dir / rel).exists() and "svg" in formats:
                dst = vdir / Path(rel).name
                shutil.copy2(project_dir / rel, dst)
                written.append(dst)
    return written


def export_visualization(v: Visualization, project_dir: Path, dest: Path) -> Path:
    """Export jedné vizualizace do zvoleného souboru (.png nebo .svg)."""
    from ..viz.render import svg_to_png
    kind = dest.suffix.lower().lstrip(".")
    if kind == "svg" and v.svg_path:
        shutil.copy2(project_dir / v.svg_path, dest)
    elif kind == "png":
        if v.png_path and (project_dir / v.png_path).exists():
            shutil.copy2(project_dir / v.png_path, dest)
        elif v.svg_path and svg_to_png((project_dir / v.svg_path).read_text(encoding="utf-8"), dest):
            pass
        else:
            raise OSError("PNG nelze vytvořit.")
    else:
        raise ValueError("Podporované formáty jsou PNG a SVG.")
    return dest
